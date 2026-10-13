"""Plain-text extraction from CV files.

The text is stored on the CV record (search, auditing, future re-analysis) and is
what Claude reads for DOCX files. PDFs are sent to Claude natively instead, so
layout and scanned pages are understood even when text extraction is poor.
"""

import logging
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.services.storage import DOCX_MIME, PDF_MIME

logger = logging.getLogger(__name__)


def extract_text(path: Path, mime_type: str) -> str | None:
    """Best-effort text extraction. Returns None if nothing could be extracted."""
    try:
        if mime_type == PDF_MIME:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif mime_type == DOCX_MIME:
            text = _docx_text(path)
        else:
            return None
    except Exception:  # corrupt / encrypted files: not fatal, Claude may still read a PDF
        logger.warning("Text extraction failed for %s", path.name, exc_info=True)
        return None
    text = text.strip()
    return text or None


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # Many CV templates put content in tables (two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
